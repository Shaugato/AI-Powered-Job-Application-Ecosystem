from __future__ import annotations

import json
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from backend.app.core.config import get_settings
from backend.app.models.entities import (
    ApplicationOutcome,
    ApplicationPlan,
    ApplicationRun,
    BaseProfile,
    CandidateMemoryProfile,
    EvidenceFragment,
    GenerationRequest,
    GenerationResult,
    MemoryEmbedding,
    OpportunityFit,
    PromptVersion,
    ResumeCompositionPlan,
    ResumeVariantScore,
    SourceAsset,
    PipelineEntityType,
    ReviewQueueReason,
    StructuredSection,
    StructuredSectionType,
    SubmissionPreview,
    TruthStatus,
)
from backend.app.schemas.corpus import (
    CorpusDocumentSummaryRead,
    CorpusIngestRequest,
    CorpusIngestResponse,
    CorpusLegacyCleanupResponse,
    LibraryFileRead,
    LibrarySectionRead,
    SourceDocumentApprovalResponse,
)
from backend.app.services.document_pipeline import DocumentPipelineService
from backend.app.services.llm import normalize_document_text
from backend.app.services.pipeline_state import PipelineStateService
from backend.app.services.review_queue import ReviewQueueService
from backend.app.services.structured_memory import StructuredMemoryService

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None


@dataclass
class ParsedDocument:
    path: Path
    body: str
    role_tags: list[str]
    document_kind: str


@dataclass
class LibraryDocumentInfo:
    key: str
    folder_name: str
    relative_path: str
    manifest_relative_path: str
    sections: list[dict[str, Any]]


class CorpusIngestionService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.settings = get_settings()
        self.memory = StructuredMemoryService(db)
        self.document_pipeline = DocumentPipelineService(db)
        self.pipeline_state = PipelineStateService(db)
        self.review_queue = ReviewQueueService(db)

    def ingest_directory(self, payload: CorpusIngestRequest, *, auto_approve: bool = True) -> CorpusIngestResponse:
        root = self._resolve_input_root(payload.path)
        if not root.exists():
            raise FileNotFoundError(f"Input path does not exist: {root}")

        warnings: list[str] = []
        source_documents: list[str] = []
        created_asset_ids: list[str] = []
        library_keys: list[str] = []
        stage_statuses: dict[str, dict[str, Any]] = {}
        created_library_manifest: dict[str, Any] = {}
        review_queue_item_ids: list[str] = []
        parsing_confidences: list[float] = []
        ambiguity_flags: list[str] = []
        document_count = 0
        section_count = 0
        prompt_count = 0
        truth_status = TruthStatus.APPROVED if auto_approve else TruthStatus.PENDING

        for file_path in sorted(root.glob(payload.source_document_glob)):
            if not file_path.is_file():
                continue
            document_kind = payload.document_kind_override or infer_document_kind(str(file_path))
            try:
                pipeline_result = self.document_pipeline.process_file(file_path, document_kind=document_kind, role_hints=payload.role_hint)
            except Exception as exc:
                warnings.append(f"Failed to parse {file_path.name}: {type(exc).__name__}: {exc}")
                continue
            parsed = ParsedDocument(
                path=file_path,
                body=pipeline_result.normalized_text,
                role_tags=sorted({*payload.role_hint, *(pipeline_result.payload.get('role_signals') or [])}),
                document_kind=document_kind,
            )
            library_info = self._materialize_library_document(parsed, pipeline_result.payload)
            if parsed.document_kind == "prompt":
                self._store_prompt(parsed)
                prompt_count += 1
            asset = self.memory.store_asset(
                source_document=str(parsed.path),
                file_name=parsed.path.name,
                asset_kind=parsed.document_kind,
                normalized_text=str(pipeline_result.payload.get("normalized_text") or parsed.body),
                library_key=library_info.key,
                library_path=library_info.relative_path,
                payload=pipeline_result.payload,
                role_tags=parsed.role_tags,
                truth_status=truth_status,
                source_metadata=pipeline_result.source_metadata,
                page_count=pipeline_result.source_metadata.get('page_count'),
                stage_records=pipeline_result.stage_records,
            )
            if pipeline_result.review_summary:
                review_item = self.review_queue.enqueue(
                    reason=ReviewQueueReason.LOW_CONFIDENCE_EXTRACTION if pipeline_result.parser_confidence < self.settings.document_review_threshold else ReviewQueueReason.LAYOUT_AMBIGUITY,
                    summary=pipeline_result.review_summary,
                    confidence=pipeline_result.parser_confidence,
                    details={
                        'ambiguity_flags': pipeline_result.ambiguity_flags,
                        'validation_result': pipeline_result.validation_result,
                        'repair_hints': pipeline_result.repair_hints,
                    },
                    asset_id=asset.id,
                    stage_name='structured_extraction',
                )
                review_queue_item_ids.append(review_item.id)
            created_asset_ids.append(asset.id)
            library_keys.append(library_info.key)
            created_library_manifest[library_info.key] = {
                'folder_name': library_info.folder_name,
                'relative_path': library_info.relative_path,
                'manifest_relative_path': library_info.manifest_relative_path,
            }
            stage_statuses[asset.id] = self.pipeline_state.stage_summary(entity_type=PipelineEntityType.ASSET, entity_id=asset.id)
            parsing_confidences.append(float(pipeline_result.parser_confidence or 0.0))
            ambiguity_flags.extend(list(pipeline_result.ambiguity_flags or []))
            document_count += 1
            source_documents.append(str(file_path))
            section_count += sum(len(section.get("files") or []) for section in library_info.sections)

        self.db.commit()
        profile = self.memory.latest_profile()
        return CorpusIngestResponse(
            ingested_documents=document_count,
            created_fragments=section_count,
            created_prompts=prompt_count,
            warnings=warnings,
            source_documents=source_documents,
            auto_approved_documents=document_count if auto_approve else 0,
            profile_id=profile.id if profile else None,
            profile_updated=bool(profile),
            created_asset_ids=created_asset_ids,
            created_section_count=section_count,
            library_keys=library_keys,
            stage_statuses=stage_statuses,
            parsing_confidence=round(sum(parsing_confidences) / max(len(parsing_confidences), 1), 4) if parsing_confidences else 0.0,
            ambiguity_flags=sorted({flag for flag in ambiguity_flags if flag}),
            review_queue_item_ids=review_queue_item_ids,
            created_library_manifest=created_library_manifest,
        )

    def _resolve_input_root(self, raw_path: str) -> Path:
        candidate = Path(str(raw_path or "")).expanduser()
        if candidate.exists():
            return candidate.resolve()

        workspace_root = Path.cwd().resolve()
        normalized = str(raw_path or "").replace("\\", "/")
        normalized_case = normalized.casefold()
        workspace_name = workspace_root.name
        marker = f"/{workspace_name}/"
        marker_case = marker.casefold()

        if marker_case in normalized_case:
            marker_index = normalized_case.index(marker_case)
            suffix = normalized[marker_index + len(marker):].strip("/")
            translated = (workspace_root / Path(suffix)).resolve()
            if translated.exists():
                return translated

        for anchor in ["private-data/", "artifacts/", "frontend/", "backend/", "infra/"]:
            anchor_case = anchor.casefold()
            if anchor_case not in normalized_case:
                continue
            anchor_index = normalized_case.index(anchor_case)
            suffix = normalized[anchor_index:].strip("/")
            translated = (workspace_root / Path(suffix)).resolve()
            if translated.exists():
                return translated

        return candidate

    def list_document_summaries(self) -> list[CorpusDocumentSummaryRead]:
        assets = self.db.query(SourceAsset).order_by(SourceAsset.created_at.desc()).all()
        sections = self.db.query(StructuredSection).order_by(StructuredSection.created_at.asc()).all()
        sections_by_asset: dict[str, list[StructuredSection]] = {}
        for section in sections:
            sections_by_asset.setdefault(section.asset_id, []).append(section)

        summaries: list[CorpusDocumentSummaryRead] = []
        for asset in assets:
            asset_sections = sections_by_asset.get(asset.id, [])
            section_counts: dict[str, int] = {}
            section_reads: list[LibrarySectionRead] = []
            for section_type in StructuredSectionType:
                typed = [section for section in asset_sections if section.section_type == section_type]
                if not typed:
                    continue
                section_counts[section_type.value] = len(typed)
                files = self._build_library_files(asset, section_type, typed)
                section_reads.append(
                    LibrarySectionRead(
                        key=section_type.value,
                        title=section_type.value.replace("_", " ").title(),
                        file_count=len(files),
                        files=files,
                    )
                )
            metadata = asset.metadata_json or {}
            truth_status = str(metadata.get("truth_status") or TruthStatus.APPROVED.value)
            pending = 1 if truth_status == TruthStatus.PENDING.value else 0
            approved = 1 if truth_status == TruthStatus.APPROVED.value else 0
            rejected = 1 if truth_status == TruthStatus.REJECTED.value else 0
            summaries.append(
                CorpusDocumentSummaryRead(
                    source_document=asset.source_document,
                    display_name=asset.file_name,
                    document_kind=asset.asset_kind.value,
                    fragment_count=len(asset_sections),
                    pending_count=pending,
                    approved_count=approved,
                    rejected_count=rejected,
                    fragment_types=sorted({section.section_type.value for section in asset_sections}),
                    role_tags=list(metadata.get("role_tags") or []),
                    skills=list({section.body_text for section in asset_sections if section.section_type == StructuredSectionType.TECHNICAL_SKILL})[:10],
                    preview_excerpt=self._preview_text(asset.normalized_text or ""),
                    library_key=asset.library_key,
                    folder_name=asset.library_key,
                    library_path=asset.library_path,
                    status=asset.parse_status,
                    parsing_confidence=float((asset.metadata_json or {}).get("parser_confidence") or 0.0),
                    ambiguity_flags=list((asset.metadata_json or {}).get("ambiguity_flags") or []),
                    review_required=bool((asset.metadata_json or {}).get("ambiguity_flags") or []) or float((asset.metadata_json or {}).get("parser_confidence") or 0.0) < self.settings.document_confidence_threshold,
                    provenance_summary={
                        "parser_version": asset.parser_version,
                        "page_count": asset.page_count,
                        "source_asset_kind": asset.asset_kind.value,
                    },
                    stage_statuses=self.pipeline_state.stage_summary(entity_type=PipelineEntityType.ASSET, entity_id=asset.id),
                    section_counts=section_counts,
                    sections=section_reads,
                    asset_id=asset.id,
                )
            )
        return summaries

    def approve_source_document(self, source_document: str, approved: bool, reviewer_notes: str | None = None) -> SourceDocumentApprovalResponse:
        fragments = self.db.query(EvidenceFragment).filter(EvidenceFragment.source_document == source_document).all()
        assets = self.db.query(SourceAsset).filter(SourceAsset.source_document == source_document).all()
        if not fragments and not assets:
            raise ValueError("Source document not found")

        new_status = TruthStatus.APPROVED if approved else TruthStatus.REJECTED
        for fragment in fragments:
            fragment.truth_status = new_status
            metadata = fragment.metadata_json or {}
            if reviewer_notes:
                metadata["reviewer_notes"] = reviewer_notes
            fragment.metadata_json = metadata
            self.db.add(fragment)
        for asset in assets:
            metadata = asset.metadata_json or {}
            metadata["truth_status"] = new_status.value
            if reviewer_notes:
                metadata["reviewer_notes"] = reviewer_notes
            asset.metadata_json = metadata
            self.db.add(asset)
        self.db.commit()
        return SourceDocumentApprovalResponse(source_document=source_document, updated_fragments=max(len(fragments), len(assets)), truth_status=new_status)

    def _parse_document(self, path: Path, role_hint: list[str], document_kind_override: str | None) -> ParsedDocument | None:
        suffix = path.suffix.lower()
        text = None
        if suffix in {".txt", ".md", ".tex"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif suffix == ".docx" and Document is not None:
            doc = Document(path)
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        elif suffix == ".pdf" and PdfReader is not None:
            reader = PdfReader(str(path))
            page_texts: list[str] = []
            for page in reader.pages:
                try:
                    extracted = page.extract_text(extraction_mode="layout") or ""
                except TypeError:
                    extracted = page.extract_text() or ""
                page_texts.append(extracted)
            text = "\n".join(page_texts)
        if not text:
            return None

        document_kind = document_kind_override or infer_document_kind(str(path))
        normalized = normalize_extracted_text(text)
        normalized = normalize_extracted_text(normalize_document_text(normalized, file_name=path.name, file_path=path, db=self.db))
        inferred_roles = sorted({*role_hint, *infer_role_tags(path.name + "\n" + normalized)})
        return ParsedDocument(path=path, body=normalized, role_tags=inferred_roles, document_kind=document_kind)

    def _materialize_library_document(self, parsed: ParsedDocument, payload: dict[str, Any]) -> LibraryDocumentInfo:
        category_dir = {
            "resume": "resumes",
            "cover_letter": "cover_letters",
            "prompt": "prompts",
        }.get(parsed.document_kind, "documents")
        prefix = {
            "resume": "resume",
            "cover_letter": "cover-letter",
            "prompt": "prompt",
        }.get(parsed.document_kind, "document")
        base_dir = self.settings.library_dir / category_dir
        base_dir.mkdir(parents=True, exist_ok=True)
        index = self._next_library_index(base_dir, prefix)
        folder_name = f"{prefix}-{index:03d}"
        folder_path = base_dir / folder_name
        folder_path.mkdir(parents=True, exist_ok=True)

        sections = self._write_library_sections(folder_path, payload)
        manifest = {
            "library_key": folder_name,
            "folder_name": folder_name,
            "display_name": parsed.path.name,
            "document_kind": parsed.document_kind,
            "source_document": str(parsed.path),
            "library_path": self._display_path(folder_path),
            "role_tags": parsed.role_tags,
            "role_signals": payload.get("role_signals") or [],
            "skills": payload.get("technical_skills") or payload.get("skills") or [],
            "status": "learned",
            "preview_excerpt": self._preview_text(str(payload.get("summary") or payload.get("normalized_text") or parsed.body)),
            "section_counts": {str(section["key"]): len(section["files"]) for section in sections},
            "sections": sections,
        }
        manifest_path = folder_path / "manifest.json"
        manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        return LibraryDocumentInfo(
            key=folder_name,
            folder_name=folder_name,
            relative_path=self._display_path(folder_path),
            manifest_relative_path=str(manifest_path),
            sections=sections,
        )

    def _write_library_sections(self, folder_path: Path, payload: dict[str, Any]) -> list[dict[str, Any]]:
        sections: list[dict[str, Any]] = []
        normalized_text = str(payload.get("normalized_text") or "").strip()
        if normalized_text:
            sections.append(self._write_text_section(folder_path, key="source", title="Source text", files=[("normalized.txt", normalized_text)]))

        identity = payload.get("identity") or {}
        identity_lines = [
            f"{label}: {value}"
            for label, value in {
                "Full name": identity.get("full_name"),
                "Email": identity.get("email"),
                "Phone": identity.get("phone"),
                "Location": identity.get("location"),
                "Website": identity.get("website"),
            }.items()
            if str(value or "").strip()
        ]
        if identity_lines:
            sections.append(self._write_text_section(folder_path, key="personal_information", title="Personal information", files=[("identity.txt", "\n".join(identity_lines))]))

        summary = str(payload.get("summary") or "").strip()
        if summary:
            sections.append(self._write_text_section(folder_path, key="summary", title="Summary", files=[("summary.txt", summary)]))

        skills = [str(item).strip() for item in payload.get("technical_skills") or payload.get("skills") or [] if str(item).strip()]
        if skills:
            sections.append(self._write_text_section(folder_path, key="technical_skills", title="Technical skills", files=[("technical_skills.txt", "\n".join(skills))]))

        certifications = [str(item).strip() for item in payload.get("certifications") or [] if str(item).strip()]
        if certifications:
            sections.append(self._write_text_section(folder_path, key="certifications", title="Certifications", files=[("certifications.txt", "\n".join(certifications))]))

        experience_entries = payload.get("experience") or []
        if experience_entries:
            files = []
            for index, entry in enumerate(experience_entries, start=1):
                title = str(entry.get("title") or f"experience-{index}").strip()
                subtitle = str(entry.get("subtitle") or "").strip()
                bullets = [str(item).strip() for item in entry.get("bullets") or [] if str(item).strip()]
                body = "\n".join([title + (f" | {subtitle}" if subtitle else ""), *[f"- {bullet}" for bullet in bullets]])
                files.append((f"{index:02d}_{slugify(title)}.txt", body))
            sections.append(self._write_text_section(folder_path, key="experience", title="Experience", files=files))

        project_entries = payload.get("projects") or []
        if project_entries:
            files = []
            for index, entry in enumerate(project_entries, start=1):
                title = str(entry.get("title") or f"project-{index}").strip()
                subtitle = str(entry.get("subtitle") or "").strip()
                bullets = [str(item).strip() for item in entry.get("bullets") or [] if str(item).strip()]
                body = "\n".join([title + (f" | {subtitle}" if subtitle else ""), *[f"- {bullet}" for bullet in bullets]])
                files.append((f"{index:02d}_{slugify(title)}.txt", body))
            sections.append(self._write_text_section(folder_path, key="projects", title="Projects", files=files))

        education = [str(item).strip() for item in payload.get("education") or [] if str(item).strip()]
        if education:
            sections.append(self._write_text_section(folder_path, key="education", title="Education", files=[("education.txt", "\n".join(education))]))

        cover_body = [str(item).strip() for item in payload.get("cover_letter_body") or [] if str(item).strip()]
        if cover_body:
            files = [(f"{index:02d}_paragraph.txt", paragraph) for index, paragraph in enumerate(cover_body, start=1)]
            sections.append(self._write_text_section(folder_path, key="cover_letter", title="Cover letter", files=files))

        prompt_logic = [str(item).strip() for item in payload.get("prompt_logic") or [] if str(item).strip()]
        if prompt_logic:
            files = [(f"{index:02d}_prompt_logic.txt", line) for index, line in enumerate(prompt_logic, start=1)]
            sections.append(self._write_text_section(folder_path, key="prompt_logic", title="Prompt logic", files=files))

        role_signals = [str(item).strip() for item in payload.get("role_signals") or [] if str(item).strip()]
        if role_signals:
            sections.append(self._write_text_section(folder_path, key="role_signals", title="Role signals", files=[("role_signals.txt", "\n".join(role_signals))]))
        return sections

    def _write_text_section(self, folder_path: Path, *, key: str, title: str, files: list[tuple[str, str]]) -> dict[str, Any]:
        section_dir = folder_path / key
        section_dir.mkdir(parents=True, exist_ok=True)
        written: list[dict[str, str | None]] = []
        for file_name, body in files:
            file_path = section_dir / file_name
            clean_body = body.strip()
            file_path.write_text(clean_body + ("\n" if clean_body else ""), encoding="utf-8")
            written.append({"name": file_name, "path": str(file_path), "preview_excerpt": self._preview_text(clean_body)})
        return {"key": key, "title": title, "files": written}

    def _next_library_index(self, base_dir: Path, prefix: str) -> int:
        highest = 0
        for item in base_dir.iterdir():
            if not item.is_dir() or not item.name.startswith(f"{prefix}-"):
                continue
            try:
                highest = max(highest, int(item.name.split("-")[-1]))
            except ValueError:
                continue
        return highest + 1

    def _store_prompt(self, parsed: ParsedDocument) -> None:
        self.db.add(PromptVersion(name=parsed.path.stem, role_tags_json=parsed.role_tags, prompt_text=parsed.body.strip(), temperature=0.2))

    def _preview_text(self, text: str) -> str:
        return re.sub(r"\s+", " ", text).strip()[:220]

    def _display_path(self, path: Path) -> str:
        try:
            return str(path.relative_to(Path.cwd()))
        except ValueError:
            return str(path)

    def _section_preview_path(self, asset: SourceAsset, section_type: StructuredSectionType, index: int) -> str:
        if not asset.library_path:
            return ""
        return str(Path(asset.library_path) / section_type.value / f"{index:02d}_{section_type.value}.txt")

    def _build_library_files(self, asset: SourceAsset, section_type: StructuredSectionType, typed: list[StructuredSection]) -> list[LibraryFileRead]:
        if section_type in {StructuredSectionType.TECHNICAL_SKILL, StructuredSectionType.CERTIFICATION, StructuredSectionType.EDUCATION}:
            lines = self._dedupe([self._clean_display_text(section.body_text) for section in typed if self._clean_display_text(section.body_text)])
            if not lines:
                return []
            file_name = {
                StructuredSectionType.TECHNICAL_SKILL: "technical_skills.txt",
                StructuredSectionType.CERTIFICATION: "certifications.txt",
                StructuredSectionType.EDUCATION: "education.txt",
            }[section_type]
            title = {
                StructuredSectionType.TECHNICAL_SKILL: "Technical skills",
                StructuredSectionType.CERTIFICATION: "Certifications",
                StructuredSectionType.EDUCATION: "Education",
            }[section_type]
            return [
                LibraryFileRead(
                    name=file_name,
                    path=None,
                    title=title,
                    lines=lines,
                    preview_excerpt=self._preview_text(", ".join(lines) if section_type != StructuredSectionType.EDUCATION else " | ".join(lines)),
                )
            ]

        files: list[LibraryFileRead] = []
        for index, section in enumerate(typed, start=1):
            title, subtitle = self._display_heading(section, index)
            lines = self._display_lines(section)
            if not lines and not title:
                continue
            files.append(
                LibraryFileRead(
                    name=f"{index:02d}_{slugify(title or section.title or section_type.value)}.txt",
                    path=None,
                    title=title,
                    subtitle=subtitle,
                    lines=lines,
                    preview_excerpt=self._preview_text(" ".join(lines) if lines else title),
                )
            )
        return files

    def _display_heading(self, section: StructuredSection, index: int) -> tuple[str | None, str | None]:
        title = self._clean_display_text(section.title)
        subtitle = self._clean_display_text(section.subtitle)
        if title and (self._looks_like_noise(title) or any(token in title.lower() for token in ["resume", "curriculum vitae", "cover letter", "prompt"])):
            title = None
        if subtitle and self._looks_like_noise(subtitle):
            subtitle = None
        if not title and subtitle:
            title, subtitle = subtitle, None
        if not title:
            title = f"{section.section_type.value.replace('_', ' ').title()} {index}"
        return title, subtitle

    def _display_lines(self, section: StructuredSection) -> list[str]:
        raw_lines = list(section.body_lines_json or []) or [section.body_text]
        lines = [self._clean_display_text(line) for line in raw_lines]
        cleaned: list[str] = []
        for line in lines:
            if not line or self._looks_like_noise(line):
                continue
            if line.casefold() in {item.casefold() for item in cleaned}:
                continue
            cleaned.append(line)
        return cleaned

    def _clean_display_text(self, value: object) -> str:
        return re.sub(r"\s+", " ", str(value or "")).strip(" -|	")

    def _looks_like_noise(self, value: str) -> bool:
        lowered = self._clean_display_text(value).lower()
        normalized = re.sub(r"[^a-z ]", "", lowered).strip()
        if not lowered:
            return True
        if lowered.startswith("private-data/") or lowered.endswith(".txt"):
            return True
        if any(token in lowered for token in ["linkedin.com", "github.com", "@", "http://", "https://"]):
            return True
        return normalized in {"experience", "education", "certifications", "projects", "technical skills", "skills", "summary", "identity", "personal information"}

    def _dedupe(self, values: list[str]) -> list[str]:
        seen: set[str] = set()
        result: list[str] = []
        for value in values:
            text = self._clean_display_text(value)
            if not text:
                continue
            key = text.casefold()
            if key in seen:
                continue
            seen.add(key)
            result.append(text)
        return result

    def cleanup_legacy_bad_uploads(self) -> CorpusLegacyCleanupResponse:
        assets = self.db.query(SourceAsset).order_by(SourceAsset.created_at.asc()).all()
        sections = self.db.query(StructuredSection).order_by(StructuredSection.created_at.asc()).all()
        sections_by_asset: dict[str, list[StructuredSection]] = {}
        for section in sections:
            sections_by_asset.setdefault(section.asset_id, []).append(section)

        bad_assets: list[SourceAsset] = []
        reasons_by_asset: dict[str, list[str]] = {}
        bad_section_ids: set[str] = set()
        library_paths: list[str] = []

        for asset in assets:
            asset_sections = sections_by_asset.get(asset.id, [])
            reasons = self._legacy_asset_reasons(asset, asset_sections)
            if not reasons:
                continue
            bad_assets.append(asset)
            reasons_by_asset[asset.id] = reasons
            bad_section_ids.update(section.id for section in asset_sections)
            if asset.library_path:
                library_paths.append(asset.library_path)

        composition_plans = self.db.query(ResumeCompositionPlan).all()
        bad_fit_ids = {fit.id for fit in self.db.query(OpportunityFit).all() if set(fit.matched_section_ids_json or []).intersection(bad_section_ids)}
        bad_composition_ids = {plan.id for plan in composition_plans if set(plan.selected_section_ids_json or []).intersection(bad_section_ids) or set(plan.excluded_section_ids_json or []).intersection(bad_section_ids)}

        generation_results = self.db.query(GenerationResult).all()
        bad_result_ids: set[str] = set()
        bad_request_ids: set[str] = set()
        for result in generation_results:
            if result.composition_plan_id and result.composition_plan_id in bad_composition_ids:
                bad_result_ids.add(result.id)
                bad_request_ids.add(result.request_id)
                continue
            if set(result.selected_section_ids_json or []).intersection(bad_section_ids):
                bad_result_ids.add(result.id)
                bad_request_ids.add(result.request_id)
                continue
            if self._generation_result_is_legacy_bad(result):
                bad_result_ids.add(result.id)
                bad_request_ids.add(result.request_id)

        bad_composition_ids.update({plan.id for plan in composition_plans if (plan.generation_request_id and plan.generation_request_id in bad_request_ids) or (plan.fit_id and plan.fit_id in bad_fit_ids)})
        bad_plan_ids = {plan.id for plan in self.db.query(ApplicationPlan).all() if plan.generation_result_id in bad_result_ids}
        bad_run_ids = {run.id for run in self.db.query(ApplicationRun).all() if run.plan_id in bad_plan_ids}

        for row in self.db.query(ApplicationOutcome).all():
            if (row.application_plan_id and row.application_plan_id in bad_plan_ids) or (row.application_run_id and row.application_run_id in bad_run_ids) or (row.generation_result_id and row.generation_result_id in bad_result_ids):
                self.db.delete(row)
        for row in self.db.query(SubmissionPreview).all():
            if row.generation_result_id in bad_result_ids or (row.application_plan_id and row.application_plan_id in bad_plan_ids):
                self.db.delete(row)
        for row in self.db.query(ResumeVariantScore).all():
            if row.generation_result_id in bad_result_ids:
                self.db.delete(row)
        for row in self.db.query(ApplicationRun).all():
            if row.id in bad_run_ids:
                self.db.delete(row)
        for row in self.db.query(ApplicationPlan).all():
            if row.id in bad_plan_ids:
                self.db.delete(row)
        self.db.flush()
        for row in self.db.query(GenerationResult).all():
            if row.id in bad_result_ids:
                self.db.delete(row)
        self.db.flush()
        for row in composition_plans:
            if row.id in bad_composition_ids:
                self.db.delete(row)
        self.db.flush()
        for row in self.db.query(GenerationRequest).all():
            if row.id in bad_request_ids:
                self.db.delete(row)
        self.db.flush()
        for row in self.db.query(OpportunityFit).all():
            if row.id in bad_fit_ids:
                self.db.delete(row)
        self.db.flush()
        for row in self.db.query(MemoryEmbedding).all():
            if row.structured_section_id in bad_section_ids:
                self.db.delete(row)
        for row in sections:
            if row.id in bad_section_ids:
                self.db.delete(row)
        self.db.flush()

        removed_source_documents = []
        removed_library_keys = []
        removed_asset_ids = []
        for asset in bad_assets:
            removed_asset_ids.append(asset.id)
            removed_source_documents.append(asset.source_document)
            if asset.library_key:
                removed_library_keys.append(asset.library_key)
            for fragment in self.db.query(EvidenceFragment).filter(EvidenceFragment.source_document == asset.source_document).all():
                self.db.delete(fragment)
            self.db.delete(asset)

        self.db.flush()
        if self.db.query(SourceAsset).count():
            self.memory._refresh_memory_profile()
        else:
            self._clear_memory_state()
        self.db.commit()

        for library_path in library_paths:
            self._delete_library_path(library_path)

        return CorpusLegacyCleanupResponse(
            removed_asset_ids=sorted(removed_asset_ids),
            removed_library_keys=sorted({key for key in removed_library_keys if key}),
            removed_source_documents=sorted({doc for doc in removed_source_documents if doc}),
            removed_generation_result_ids=sorted(bad_result_ids),
            removed_application_plan_ids=sorted(bad_plan_ids),
            removed_application_run_ids=sorted(bad_run_ids),
            removed_section_ids=sorted(bad_section_ids),
            removed_request_ids=sorted(bad_request_ids),
            removed_fit_ids=sorted(bad_fit_ids),
            reasons_by_asset=reasons_by_asset,
        )

    def _legacy_asset_reasons(self, asset: SourceAsset, sections: list[StructuredSection]) -> list[str]:
        parser_version = str(asset.parser_version or "")
        anomaly_reasons: list[str] = []
        source_document = str(asset.source_document or "").lower()
        if any(token in source_document for token in ["debug-upload", "sample-data", "resume-devops.md"]):
            anomaly_reasons.append("debug_fixture_upload")
        if not sections:
            anomaly_reasons.append("missing_sections")
        else:
            for section in sections:
                anomaly_reasons.extend(self._legacy_section_reasons(asset, section))
        anomaly_reasons = self._dedupe(anomaly_reasons)
        if not anomaly_reasons:
            return []
        if parser_version != "structured-memory-v2":
            anomaly_reasons.insert(0, f"parser_version:{parser_version or 'unknown'}")
        return self._dedupe(anomaly_reasons)

    def _legacy_section_reasons(self, asset: SourceAsset, section: StructuredSection) -> list[str]:
        reasons: list[str] = []
        section_type = section.section_type
        title = self._clean_display_text(section.title)
        subtitle = self._clean_display_text(section.subtitle)
        body_lines = [self._clean_display_text(line) for line in list(section.body_lines_json or [])]
        combined_lines = [line for line in [title, subtitle, *body_lines, self._clean_display_text(section.body_text)] if line]

        if section_type in {StructuredSectionType.EXPERIENCE, StructuredSectionType.PROJECT}:
            if self._is_container_title(title) or self._is_container_title(subtitle):
                reasons.append(f"{section_type.value}:container_title")
            if any(self._contains_contact_or_path(line) for line in combined_lines):
                reasons.append(f"{section_type.value}:contact_or_path")
            if any(self._looks_like_heading(line) for line in body_lines):
                reasons.append(f"{section_type.value}:heading_leak")

        if section_type == StructuredSectionType.TECHNICAL_SKILL:
            text = self._clean_display_text(section.body_text)
            if self._contains_contact_or_path(text):
                reasons.append("technical_skill:contact_or_path")
            if len(text.split()) > 8:
                reasons.append("technical_skill:not_atomic")

        if section_type in {StructuredSectionType.SUMMARY, StructuredSectionType.CERTIFICATION, StructuredSectionType.EDUCATION}:
            if any(self._contains_contact_or_path(line) for line in combined_lines):
                reasons.append(f"{section_type.value}:contact_or_path")

        if section_type == StructuredSectionType.IDENTITY and not any(part.strip() for part in body_lines):
            reasons.append("identity:empty")

        if asset.file_name and title and self._section_title_matches_file(asset.file_name, title):
            reasons.append(f"{section_type.value}:file_name_leak")

        return reasons

    def _generation_result_is_legacy_bad(self, result: GenerationResult) -> bool:
        if self._contains_contact_or_path(result.grounding_notes or "") or re.search(r"invalid schema|badrequesterror|invalid_json_schema|fallback to local heuristic", str(result.grounding_notes or ""), re.I):
            return True
        resume = result.resume_json or {}
        for key in ["summary"]:
            if self._contains_contact_or_path(str(resume.get(key) or "")):
                return True
        for key in ["skills", "certifications", "education"]:
            for value in list(resume.get(key) or []):
                if self._contains_contact_or_path(str(value)):
                    return True
        for key in ["experience", "projects"]:
            for entry in list(resume.get(key) or []):
                title = self._clean_display_text((entry or {}).get("title"))
                subtitle = self._clean_display_text((entry or {}).get("subtitle"))
                bullets = [self._clean_display_text(item) for item in ((entry or {}).get("bullets") or [])]
                if self._is_container_title(title) or self._is_container_title(subtitle):
                    return True
                if any(self._contains_contact_or_path(line) or self._looks_like_heading(line) for line in [title, subtitle, *bullets] if line):
                    return True
        return False

    def _contains_contact_or_path(self, value: str) -> bool:
        lowered = self._clean_display_text(value).lower()
        if not lowered:
            return False
        if lowered.startswith("private-data/") or lowered.startswith("source: private-data/") or lowered.endswith(".txt"):
            return True
        return any(token in lowered for token in ["linkedin.com", "github.com", "http://", "https://", "@", "\\private-data\\"])

    def _looks_like_heading(self, value: str) -> bool:
        lowered = re.sub(r"[^a-z ]", "", self._clean_display_text(value).lower()).strip()
        return lowered in {"experience", "education", "certifications", "projects", "technical skills", "skills", "summary", "identity", "personal information"}

    def _is_container_title(self, value: str | None) -> bool:
        lowered = self._clean_display_text(value).lower()
        if not lowered:
            return False
        if lowered in {"resume", "resumes", "cover letter", "prompt", "resume batch", "resume-batch", "debug upload", "debug-upload"}:
            return True
        return lowered.endswith(" resume") or lowered.endswith(" cv") or lowered.startswith("resume ")

    def _section_title_matches_file(self, file_name: str, title: str) -> bool:
        file_stem = slugify(Path(file_name).stem)
        title_slug = slugify(title)
        return bool(file_stem and title_slug and (title_slug == file_stem or title_slug in file_stem or file_stem in title_slug))

    def _delete_library_path(self, relative_or_absolute: str) -> None:
        if not relative_or_absolute:
            return
        candidate = Path(relative_or_absolute)
        if not candidate.is_absolute():
            candidate = (Path.cwd() / candidate).resolve()
        else:
            candidate = candidate.resolve()
        library_root = self.settings.library_dir.resolve()
        try:
            candidate.relative_to(library_root)
        except ValueError:
            return
        if candidate.exists() and candidate.is_dir():
            shutil.rmtree(candidate, ignore_errors=True)

    def _clear_memory_state(self) -> None:
        for profile in self.db.query(CandidateMemoryProfile).all():
            self.db.delete(profile)
        for profile in self.db.query(BaseProfile).all():
            profile.skills_json = []
            profile.certifications_json = []
            profile.experience_json = []
            profile.projects_json = []
            profile.education_json = []
            metadata = profile.metadata_json or {}
            metadata["memory_source_documents"] = []
            metadata["autofill_mode"] = "structured_memory"
            profile.metadata_json = metadata
            self.db.add(profile)



def normalize_extracted_text(text: str) -> str:
    replacements = {
        "\u00a0": " ",
        "\u200b": "",
        "\uf0b7": "- ",
        "\u2022": "- ",
        "\u2023": "- ",
        "\u25e6": "- ",
        "\ufffd": " ",
        "\uf0d8": " ",
        "\u2013": " - ",
        "\u2014": " - ",
        "\u00b7": " | ",
        "\u2197": " ",
    }
    normalized = unicodedata.normalize("NFKC", text or "")
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    normalized = re.sub(r"(?i)linkedinlinkedin", "linkedin ", normalized)
    normalized = re.sub(r"(?i)githubgithub", "github ", normalized)
    normalized = re.sub(r"[ \t]{2,}", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    lines = [re.sub(r"\s+", " ", line).strip(" |\t") for line in normalized.splitlines()]
    cleaned_lines: list[str] = []
    previous = None
    for line in lines:
        if not line:
            if previous != "":
                cleaned_lines.append("")
            previous = ""
            continue
        if line == previous:
            continue
        cleaned_lines.append(line)
        previous = line
    return "\n".join(cleaned_lines).strip()


def infer_role_tags(text: str) -> list[str]:
    text_lower = text.lower()
    mapping = {
        "DevOps": ["devops", "terraform", "ci/cd", "sre"],
        "Cloud": ["aws", "azure", "gcp", "cloud"],
        "Networking": ["network", "routing", "switching", "ccna"],
        "Security": ["security", "siem", "soc", "iam"],
        "IT Support": ["service desk", "desktop", "support"],
        "Infrastructure": ["vmware", "windows server", "linux", "infrastructure"],
    }
    return [role for role, needles in mapping.items() if any(needle in text_lower for needle in needles)]


def infer_document_kind(source_document: str) -> str:
    text = source_document.lower()
    if "cover" in text:
        return "cover_letter"
    if "prompt" in text:
        return "prompt"
    if any(token in text for token in ["resume", "cv"]):
        return "resume"
    return "document"


def slugify(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", str(value or "").strip().lower()).strip("-")
    return cleaned or "entry"
